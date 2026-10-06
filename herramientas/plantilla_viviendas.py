"""Genera la plantilla Excel de la tabla de viviendas y tipologías.

La descarga la promotora desde su portal (requisito «Tabla de viviendas y
tipologías»), la rellena y la sube. Uso:

    python3 herramientas/plantilla_viviendas.py

Escribe portal/public/plantillas/tabla-viviendas.xlsx.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DESTINO = Path(__file__).resolve().parent.parent / 'portal' / 'public' / 'plantillas' / 'tabla-viviendas.xlsx'

COLUMNAS = [
    ('Vivienda', 'Identificador de la vivienda, tal como se comercializa (p. ej. 1ºA, P1-B, Parcela 12)', 14),
    ('Portal o bloque', 'Solo en plurifamiliares', 14),
    ('Planta', 'Baja, 1, 2… o Ático', 10),
    ('Letra o puerta', '', 12),
    ('Tipología', 'Nombre de la tipología (p. ej. Tipo A)', 14),
    ('Variante', 'Espejo, bajo con jardín, ático… (vacío si no tiene)', 18),
    ('Dormitorios', '', 12),
    ('Baños', '', 8),
    ('Superficie útil (m²)', '', 14),
    ('Superficie construida (m²)', 'Sin zonas comunes', 16),
    ('Superficie construida con comunes (m²)', '', 18),
    ('Terraza (m²)', '', 12),
    ('Jardín o patio (m²)', '', 14),
    ('Garaje', 'Número de plaza(s), si va vinculada', 12),
    ('Trastero', 'Número de trastero, si va vinculado', 12),
    ('Orientación', 'N, S, E, O, NE…', 12),
    ('Observaciones', '', 30),
]

EJEMPLO = ['1ºA', '1', '1', 'A', 'Tipo A', '', 3, 2, 85.40, 102.15, 118.60, 12.30, '', 'P-14', 'T-3', 'S', 'Ejemplo: borre esta fila']


def main() -> None:
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Viviendas'
    cabecera = PatternFill('solid', fgColor='2F5D50')
    for i, (nombre, ayuda, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=nombre)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = cabecera
        celda.alignment = Alignment(wrap_text=True, vertical='center')
        ayuda_celda = hoja.cell(row=2, column=i, value=ayuda)
        ayuda_celda.font = Font(italic=True, color='8B867D', size=9)
        ayuda_celda.alignment = Alignment(wrap_text=True, vertical='top')
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for i, valor in enumerate(EJEMPLO, start=1):
        hoja.cell(row=3, column=i, value=valor).font = Font(color='8B867D')
    hoja.row_dimensions[1].height = 32
    hoja.row_dimensions[2].height = 40
    hoja.freeze_panes = 'B3'

    notas = libro.create_sheet('Instrucciones')
    for fila, texto in enumerate([
        'Tabla de viviendas y tipologías',
        '',
        '1. Una fila por vivienda, empezando en la fila 3 de la hoja «Viviendas» (la fila 3 es un ejemplo: bórrela).',
        '2. No cambie el nombre ni el orden de las columnas.',
        '3. Las superficies, en metros cuadrados con decimales (85,40).',
        '4. Si un dato no aplica, deje la celda vacía.',
        '5. Guarde el archivo y súbalo en el portal, en «Tabla de viviendas y tipologías».',
    ], start=1):
        notas.cell(row=fila, column=1, value=texto).font = Font(bold=fila == 1, size=13 if fila == 1 else 11)
    notas.column_dimensions['A'].width = 110

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    libro.save(DESTINO)
    print(f'Plantilla escrita en {DESTINO}')


if __name__ == '__main__':
    main()
